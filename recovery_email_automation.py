"""
Seta o recovery email (contact@wolfscale.com) de todos os usuarios de cada
dominio pendente na "Mailbox Recovery Project.xlsx", logando como o admin
de cada workspace num navegador visivel (Playwright), janela isolada por
dominio (equivalente a aba anonima).

Uso:
    python recovery_email_automation.py           # processa 20 dominios (default)
    python recovery_email_automation.py 5          # processa so 5 (bom pro primeiro teste)

Progresso e salvo na propria planilha (coluna done) apos cada dominio,
entao rodar de novo continua de onde parou.
"""

import sys
import time
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

XLSX_PATH = Path(__file__).resolve().parent.parent / "Mailbox Recovery Project.xlsx"
LOCK_PATH = XLSX_PATH.parent / f"~$Mailbox Recovery Project.xlsx"
FAILURES_LOG = Path(__file__).resolve().parent / "failures.log"
SHEET_NAME = "mailboxes_export"
RECOVERY_EMAIL = "contact@wolfscale.com"
DEFAULT_BATCH_SIZE = 20

COL_CUSTOMER = 1
COL_EMAIL = 2
COL_PASSWORD = 3
COL_BACKUP = 4
COL_DOMAIN = 5
COL_DONE = 6
COL_NOTES = 7

PASSWORD_LOOKUP_SHEET = "Sheet1"
PASSWORD_LOOKUP_EMAIL_COL = 1  # A
PASSWORD_LOOKUP_PASSWORD_COL = 2  # B


def build_password_lookup(wb):
    # coluna C da aba principal e uma formula VLOOKUP pra Sheet1, e o
    # openpyxl nao recalcula formulas (o cache do valor calculado se
    # perde ao salvar por aqui). A fonte real e sempre a Sheet1.
    ws2 = wb[PASSWORD_LOOKUP_SHEET]
    lookup = {}
    for row in ws2.iter_rows(values_only=True):
        email = row[PASSWORD_LOOKUP_EMAIL_COL - 1]
        password = row[PASSWORD_LOOKUP_PASSWORD_COL - 1]
        if email:
            lookup[email] = password
    return lookup


def is_pending(done):
    # a coluna done pode ser um booleano nativo do Excel (False) ou
    # texto "FALSE" (depois que a planilha foi convertida pra texto em
    # ingles), entao aceita as duas formas.
    if done is False:
        return True
    if isinstance(done, str) and done.strip().upper() == "FALSE":
        return True
    return False


def load_target_domains(wb, password_lookup):
    ws = wb[SHEET_NAME]
    rows = []
    for row_idx in range(2, ws.max_row + 1):
        email = ws.cell(row_idx, COL_EMAIL).value
        domain = ws.cell(row_idx, COL_DOMAIN).value
        done = ws.cell(row_idx, COL_DONE).value
        notes = ws.cell(row_idx, COL_NOTES).value
        if is_pending(done) and notes is None:
            password = password_lookup.get(email)
            rows.append({"row": row_idx, "email": email, "password": password, "domain": domain})

    by_domain = {}
    for r in rows:
        by_domain.setdefault(r["domain"], []).append(r)

    domains = []
    for domain, users in by_domain.items():
        no_dot = [u for u in users if "." not in u["email"].split("@")[0]]
        admin = None
        for u in no_dot:
            local = u["email"].split("@")[0]
            is_prefixed_by_another = any(
                o["email"] != u["email"] and local.startswith(o["email"].split("@")[0])
                for o in no_dot
            )
            if not is_prefixed_by_another:
                admin = u
                break
        if admin is None:
            # nao deveria acontecer: dominios sem admin claro ja foram
            # marcados "off" ou "find who is the admin" na planilha
            continue
        if not admin["password"]:
            print(f"  aviso: sem senha encontrada na Sheet1 para {admin['email']}, marcando {domain} e pulando")
            for u in users:
                ws.cell(u["row"], COL_NOTES).value = "password not found"
            continue
        domains.append({"domain": domain, "admin": admin, "users": users})
    return domains


def wait_for_manual_checkpoint(page, reason):
    # NEM TODA CONTA vai cair aqui. A maioria loga direto sem parar.
    # Isso so pausa quando O GOOGLE especificamente pede algo extra
    # (codigo de verificacao, captcha, confirmar recovery, etc) pra
    # ESSA conta. Assim que voce resolver na janela do navegador e
    # apertar ENTER aqui no terminal, o script CONTINUA exatamente de
    # onde parou (nao reinicia, nao pula o dominio, nao repete passos
    # anteriores). O input() abaixo trava o script ate voce apertar ENTER.
    try:
        page.bring_to_front()
    except Exception:
        pass
    print(f"\n>>> PAUSADO: {reason}")
    print(">>> So acontece nas contas que o Google decide pedir verificacao extra.")
    print(">>> NAO TEM PRESSA E NAO TEM LIMITE DE TEMPO. Pode esperar o codigo chegar")
    print(">>> no email, demorar o quanto precisar. NAO aperte Ctrl+C, isso interrompe")
    print(">>> o script no meio (nao cancela so essa espera). So aperte ENTER quando resolver.")
    input(">>> Resolva na janela do navegador e pressione ENTER aqui pra continuar de onde parou...\n")


def looks_like_verification_prompt(page):
    indicators = [
        "2-Step Verification",
        "Verify it's you",
        "Confirm your recovery",
        "unusual activity",
    ]
    try:
        content = page.content()
    except Exception:
        return False
    return any(text.lower() in content.lower() for text in indicators)


def login_as_admin(context, admin):
    page = context.new_page()
    page.goto(
        "https://accounts.google.com/ServiceLogin?service=accountsettings&continue=https://admin.google.com/&hl=en"
    )
    page.wait_for_selector("#identifierId", timeout=30000)
    page.fill("#identifierId", admin["email"])
    page.click("#identifierNext button")
    page.wait_for_timeout(2000)

    if looks_like_verification_prompt(page):
        wait_for_manual_checkpoint(page, "Tela inesperada apos digitar o email (pode ser captcha/verificacao).")

    page.wait_for_selector('input[type="password"]', timeout=30000)
    page.fill('input[type="password"]', admin["password"])
    page.click("#passwordNext button")
    print("  verificando se o Google vai pedir confirmacao extra, aguarde uns segundos (normal)...")
    page.wait_for_timeout(2500)

    if looks_like_verification_prompt(page):
        wait_for_manual_checkpoint(page, "2FA / verificacao adicional detectada.")

    print("  carregando o admin console, pode levar ate 20s (normal, nao e trava)...")
    try:
        page.wait_for_url("**/admin.google.com/**", timeout=20000)
    except Exception:
        wait_for_manual_checkpoint(
            page, "Nao chegou no admin.google.com ainda. Confirme manualmente e pressione ENTER quando estiver la."
        )

    # depois de logar, o Google ainda faz alguns redirecionamentos internos
    # (rapt=..., pli=1, etc) antes de estabilizar. Da tempo pra isso acabar
    # antes de navegar de novo, senao uma navegacao atropela a outra.
    try:
        page.wait_for_load_state("load", timeout=15000)
    except Exception:
        pass
    page.wait_for_timeout(3000)

    return page


def goto_with_retry(page, url, attempts=3):
    last_error = None
    for _ in range(attempts):
        try:
            page.goto(url)
            return
        except Exception as e:
            last_error = e
            page.wait_for_timeout(2000)
    raise last_error


def goto_users_list(page):
    goto_with_retry(page, "https://admin.google.com/ac/users?hl=en")
    # "text=Users" aparece varias vezes na tela (menu, titulo, etc), o
    # que confunde o wait. Confirma pela URL, que e unica, em vez disso.
    page.wait_for_url("**/ac/users**", timeout=20000)
    page.wait_for_timeout(1500)


def set_recovery_email_for_user(page, user_email):
    # a lista de Users ja mostra todo mundo do dominio direto na tela
    # (dominios pequenos, poucos usuarios), entao clica direto na linha
    # em vez de tentar buscar (a caixa de busca do topo e um "Ask me
    # anything" de IA, nao um filtro de lista de usuarios).
    #
    # varios usuarios podem ter o MESMO nome (ex: "Joe Davy" nas 3
    # variantes de banzaibrand.com), entao nao da pra clicar por nome.
    # em vez disso, acha a LINHA da tabela que contem o email exato, e
    # dentro dela clica no link real (<a href="./ac/users/{id}">).
    row = page.locator("tr", has_text=user_email)
    row.locator('a[href*="/ac/users/"]').first.click()
    page.wait_for_timeout(2000)

    # a palavra "Security" tambem existe no menu lateral (configuracoes
    # gerais da organizacao), entao get_by_text pegaria os dois. Usa role
    # "tab" pra pegar so a aba dentro do painel do usuario.
    try:
        page.get_by_role("tab", name="Security", exact=True).click()
    except Exception:
        page.get_by_text("Security", exact=True).last.click()
    page.wait_for_timeout(1500)

    # "Recovery information" mostra "Add a recovery email" como um link
    # (nao um campo pronto), so vira campo de texto depois de clicar nele.
    page.get_by_text("Add a recovery email", exact=True).first.click()
    page.wait_for_timeout(1000)

    # o modal tem 2 campos (Email e Phone), tem que mirar especificamente
    # no de Email, senao o texto vai pro campo errado (Phone).
    recovery_field = page.get_by_role("textbox", name="Email", exact=False).first
    recovery_field.click()
    recovery_field.fill(RECOVERY_EMAIL)
    page.wait_for_timeout(500)
    page.get_by_role("button", name="Save").first.click()
    page.wait_for_timeout(1500)

    # volta pra lista de usuarios pra processar o proximo
    goto_users_list(page)


DEBUG_DIR = Path(__file__).resolve().parent / "debug_screenshots"


def process_domain(browser, entry, wb):
    domain = entry["domain"]
    admin = entry["admin"]
    print(f"\n--- {domain} (admin: {admin['email']}, {len(entry['users'])} usuario(s)) ---")

    context = browser.new_context()
    page = None
    try:
        page = login_as_admin(context, admin)
        goto_users_list(page)
        for u in entry["users"]:
            print(f"  recovery email -> {u['email']}")
            set_recovery_email_for_user(page, u["email"])
        for u in entry["users"]:
            wb[SHEET_NAME].cell(u["row"], COL_DONE).value = "TRUE"
        return True, None
    except Exception as e:
        print(f"  ERRO: {e}")
        # notes NAO e marcado aqui de proposito: assim o dominio fica
        # elegivel pra tentar de novo no proximo run sem precisar limpar
        # nada manualmente. Falhas ficam registradas so no log abaixo.
        with open(FAILURES_LOG, "a", encoding="utf-8") as f:
            f.write(f"{domain} | admin={admin['email']} | {e}\n")
        if page is not None:
            DEBUG_DIR.mkdir(exist_ok=True)
            safe_name = domain.replace(".", "_")
            try:
                page.screenshot(path=str(DEBUG_DIR / f"{safe_name}.png"), full_page=True)
                (DEBUG_DIR / f"{safe_name}.html").write_text(page.content(), encoding="utf-8")
                print(f"  screenshot salva em debug_screenshots/{safe_name}.png")
            except Exception as screenshot_error:
                print(f"  (nao consegui salvar screenshot: {screenshot_error})")
        return False, str(e)
    finally:
        context.close()


def main():
    if LOCK_PATH.exists():
        print(f"O arquivo '{XLSX_PATH.name}' esta aberto no Excel. Feche e rode de novo.")
        sys.exit(1)

    batch_size = int(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_BATCH_SIZE

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=False)
    password_lookup = build_password_lookup(wb)
    domains = load_target_domains(wb, password_lookup)
    wb.save(XLSX_PATH)  # persiste os "password not found" marcados agora
    batch = domains[:batch_size]

    print(f"Dominios pendentes no total: {len(domains)}")
    print(f"Processando {len(batch)} agora (batch_size={batch_size})\n")

    results = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        for entry in batch:
            ok, err = process_domain(browser, entry, wb)
            results.append((entry["domain"], ok, err))
            wb.save(XLSX_PATH)
        browser.close()

    sucesso = sum(1 for _, ok, _ in results if ok)
    falha = len(results) - sucesso

    print("\n=== Resumo do batch ===")
    for domain, ok, err in results:
        status = "OK" if ok else f"FALHOU: {err}"
        print(f"  {domain}: {status}")

    print(f"\n{sucesso} concluido(s) com sucesso, {falha} falhou(aram) neste batch.")
    print(f"Dominios ainda pendentes no total (incluindo falhas, tentaveis de novo): {len(domains) - sucesso}")


if __name__ == "__main__":
    main()
